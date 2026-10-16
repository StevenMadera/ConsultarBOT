"""Las filas pegadas desde Excel (cédula, fecha de expedición y nombre), qué dato está mal y el Excel de salida."""
import csv
import io
import itertools
import math
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import nombres
from comun import NOMBRE, SALIDAS

TIPOS = ("CC", "TI", "CE", "PEP", "PPT", "NIT")
MESES = {"ENE": 1, "JAN": 1, "FEB": 2, "MAR": 3, "ABR": 4, "APR": 4, "MAY": 5, "JUN": 6, "JUL": 7, "AGO": 8,
         "AUG": 8, "SEP": 9, "SET": 9, "OCT": 10, "NOV": 11, "DIC": 12, "DEC": 12}
TITULO_RE = re.compile(r"CEDULA|DOCUMENTO|IDENTIFICACION|NOMBRE|APELLIDO|EXPEDICION|FECHA|\bCC\b|\bEXP\b")
PALABRAS_SUELTAS = {"DE", "DEL", "LA", "LAS", "LOS", "Y", "VDA", "VIUDA"}
SITIOS = {"libreta": "la libreta militar", "simit": "el SIMIT", "rnmc": "el RNMC (medidas correctivas)",
          "lote": "la fila"}

ROJO = PatternFill("solid", fgColor="FFC7CE")
AMARILLO = PatternFill("solid", fgColor="FFEB9C")
GRIS = PatternFill("solid", fgColor="D9D9D9")
LINEA = Side(style="thin", color="A6A6A6")
BORDE = Border(left=LINEA, right=LINEA, top=LINEA, bottom=LINEA)


def normalizar(texto) -> str:
    """Sin tildes, en mayúsculas y solo letras y números separados por un espacio."""
    t = "".join(c for c in unicodedata.normalize("NFD", str(texto or "")) if unicodedata.category(c) != "Mn")
    return " ".join(re.sub(r"[^A-Z0-9Ñ]+", " ", t.upper()).split())


def leer_fecha(texto: str) -> date:
    """25/06/2020, 25-06-2020, 2020-06-25, 25/06/20, 25-jun-2020, 25 de junio de 2020 o el número de serie de Excel."""
    t = " ".join(str(texto).split())
    t = re.sub(r"\s+\d{1,2}:\d{2}(?::\d{2})?(?:\s*[ap]\.?\s*m\.?)?$", "", t, flags=re.IGNORECASE)
    if re.fullmatch(r"\d{5}", t) and 10000 < int(t) < 80000:
        return date(1899, 12, 30) + timedelta(days=int(t))
    if m := re.fullmatch(r"(\d{4})[/.-](\d{1,2})[/.-](\d{1,2})", t):
        a, mes, d = int(m[1]), int(m[2]), int(m[3])
    elif m := re.fullmatch(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{2}|\d{4})", t):
        d, mes, a = int(m[1]), int(m[2]), int(m[3])
    elif m := re.fullmatch(r"(\d{1,2})[\s/.-]+(?:de\s+)?([A-Za-zñÑáéíóúÁÉÍÓÚ]{3,})\.?[\s/.-]+(?:de\s+)?(\d{2}|\d{4})", t,
                           flags=re.IGNORECASE):
        d, a = int(m[1]), int(m[3])
        mes = MESES.get(normalizar(m[2])[:3])
        if mes is None:
            raise ValueError(texto)
    else:
        raise ValueError(texto)
    if a < 100:
        a += 2000 if a <= date.today().year % 100 else 1900
    return date(a, mes, d)


def es_fecha(texto: str) -> bool:
    try:
        leer_fecha(texto)
        return True
    except ValueError:
        return False


def documento(texto: str) -> tuple[str, str] | None:
    """("CC", "1234567890") a partir de "1234567890", "1.234.567.890" o "TI 1234567890"."""
    m = re.fullmatch(r"(?:([A-Za-z.]{2,4})\s*)?([\d][\d.,\s]*)", str(texto).strip())
    if not m:
        return None
    tipo = (m[1] or "CC").replace(".", "").upper()
    numero = re.sub(r"\D", "", m[2])
    if tipo not in TIPOS or not 4 <= len(numero) <= 11:
        return None
    return tipo, numero


def es_nombre(texto: str) -> bool:
    texto = " ".join(texto.split())
    return len(texto.split()) >= 2 and bool(re.fullmatch(r"[A-Za-zÀ-ÿ'´.\- ]+", texto))


def deducir_sexo(nombre: str) -> str:
    """Por los nombres de pila conocidos; si empatan, decide el primero. "" si no reconoce ninguno."""
    votos, primero = {"H": 0, "M": 0}, ""
    for palabra in normalizar(nombre).split():
        sexo = "M" if palabra in nombres.MUJER else "H" if palabra in nombres.HOMBRE else ""
        if sexo:
            votos[sexo] += 1
            primero = primero or sexo
    if votos["H"] == votos["M"]:
        return primero
    return "H" if votos["H"] > votos["M"] else "M"


def mismo_nombre(a: str, b: str) -> bool:
    """La mayoría de las palabras del nombre más corto están en el otro, con tolerancia a errores de escritura."""
    pa = [p for p in normalizar(a).split() if p not in PALABRAS_SUELTAS]
    pb = [p for p in normalizar(b).split() if p not in PALABRAS_SUELTAS]
    if not pa or not pb:
        return True
    corto, largo = sorted((pa, pb), key=len)
    iguales = sum(any(SequenceMatcher(None, p, q).ratio() >= 0.8 for q in largo) for p in corto)
    return iguales >= math.ceil(len(corto) * 0.6)


@dataclass
class Fila:
    numero: int
    celdas: list[str]
    # Fila en blanco entre las pegadas: se conserva para que las observaciones queden alineadas.
    vacia: bool = False
    tipo: str = "CC"
    cedula: str = ""
    nombre: str = ""
    expedicion: date | None = None
    sexo: str = ""
    # Datos mal que impiden consultar: "REVISAR: ..." en rojo.
    errores: list[str] = field(default_factory=list)
    # La observación vale, pero algo no cuadra (el nombre): en rojo, con una nota.
    avisos: list[str] = field(default_factory=list)
    # Sitio que no respondió -> motivo (en amarillo).
    fallas: dict[str, str] = field(default_factory=dict)
    libreta: object = None
    simit: object = None
    anotaciones: list | None = None
    observacion: str = ""

    @property
    def consultable(self) -> bool:
        """Con cédula y fecha de expedición válidas; si no, no se consulta y se marca para corregir."""
        return not self.vacia and bool(self.cedula) and self.expedicion is not None and not self.errores


@dataclass
class Tabla:
    filas: list[Fila]


def separar(texto: str) -> list[list[str]]:
    """Celdas separadas por tabulaciones, como las copia Excel (las que tienen saltos de línea van entre comillas).
    Las filas en blanco del medio se conservan."""
    filas = [[c.strip() for c in f] for f in csv.reader(io.StringIO(texto.strip("\r\n")), delimiter="\t")]
    while filas and not any(filas[0]):
        filas.pop(0)
    while filas and not any(filas[-1]):
        filas.pop()
    if not filas:
        return []
    ancho = max(len(f) for f in filas)
    filas = [f + [""] * (ancho - len(f)) for f in filas]
    # Columnas vacías a la derecha, si se copió de más.
    while ancho > 1 and not any(f[ancho - 1] for f in filas):
        ancho -= 1
    return [f[:ancho] for f in filas]


def orden_columnas(datos: list[list[str]]) -> tuple[int, int, int]:
    """Columnas de la cédula, la fecha de expedición y el nombre: en ese orden, salvo que el contenido diga otro."""
    pruebas = (documento, es_fecha, es_nombre)

    def puntaje(orden):
        total = 0.0
        for columna, prueba in zip(orden, pruebas):
            llenos = [f[columna] for f in datos if f[columna]]
            total += sum(1 for v in llenos if prueba(v)) / len(llenos) if llenos else 0
        return total

    # max() se queda con el primero de los empatados: el orden pedido (0, 1, 2).
    return max(itertools.permutations(range(3)), key=puntaje)


def armar(texto: str) -> Tabla:
    """Las filas a partir del texto copiado. ValueError si no llegan exactamente las 3 columnas."""
    crudas = separar(texto)
    if not crudas:
        raise ValueError("No hay filas: copia en Excel la cédula, la fecha de expedición y el nombre.")
    ancho = len(crudas[0])
    if ancho != 3:
        raise ValueError(f"Llegaron {ancho} columna{'s' if ancho != 1 else ''}: copia en Excel solo 3, la cédula, "
                         "la fecha de expedición y el nombre (en ese orden).")
    # La fila de títulos, si se copió, no se consulta.
    if not any(documento(c) or es_fecha(c) for c in crudas[0]) and any(TITULO_RE.search(normalizar(c)) for c in crudas[0]):
        crudas = crudas[1:]
        if not crudas:
            raise ValueError("Solo llegó la fila de títulos: copia también las filas de las personas.")
    cedula, expedicion, nombre = orden_columnas([f for f in crudas if any(f)])
    return Tabla([validar_fila(Fila(n, [f[cedula], f[expedicion], f[nombre]]))
                  for n, f in enumerate(crudas, start=1)])


def validar_fila(fila: Fila) -> Fila:
    cedula, expedicion, nombre = fila.celdas
    if not any(fila.celdas):
        fila.vacia = True
        return fila
    doc = documento(cedula)
    if doc is None:
        fila.errores.append(f"la cédula «{cedula}» no es válida" if cedula else "falta la cédula")
    else:
        fila.tipo, fila.cedula = doc
    if not expedicion:
        fila.errores.append("falta la fecha de expedición")
    else:
        try:
            fila.expedicion = leer_fecha(expedicion)
        except ValueError:
            fila.errores.append(f"la fecha de expedición «{expedicion}» no es una fecha válida")
        else:
            if fila.expedicion > date.today():
                fila.errores.append(f"la fecha de expedición {fila.expedicion:%d/%m/%Y} es posterior a hoy")
                fila.expedicion = None
    fila.nombre = " ".join(nombre.split())
    if not fila.nombre:
        fila.errores.append("falta el nombre")
    fila.sexo = deducir_sexo(fila.nombre)
    return fila


def texto_revisar(fila: Fila) -> str:
    motivos = fila.errores + [f"no se pudo consultar {SITIOS[sitio]} ({motivo})" for sitio, motivo in fila.fallas.items()]
    texto = "; ".join(motivos)
    return f"REVISAR: {texto[:1].upper()}{texto[1:]}."


def texto_final(fila: Fila) -> str:
    """Lo que va en OBSERVACIÓN: el resultado, o qué revisar si un dato está mal o un sitio no respondió."""
    return texto_revisar(fila) if fila.errores or fila.fallas else fila.observacion


def guardar_excel(tabla: Tabla, ruta: Path | None = None) -> Path:
    """CÉDULA y OBSERVACIÓN, una fila por cada fila pegada y en el mismo orden. Roja: dato por corregir (o el
    nombre no coincide: va en una nota); amarilla: un sitio no respondió."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Observaciones"
    hoja.append(["CÉDULA", "OBSERVACIÓN"])
    for celda in hoja[1]:
        celda.font = Font(bold=True)
        celda.fill = GRIS
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = BORDE
    for n, fila in enumerate(tabla.filas, start=2):
        # Como número, para que se pueda cruzar (BUSCARV) con la cédula del Excel de origen.
        cedula = hoja.cell(row=n, column=1, value=int(fila.cedula) if fila.cedula else fila.celdas[0] or None)
        cedula.number_format = "0"
        observacion = hoja.cell(row=n, column=2, value=texto_final(fila) or None)
        for celda in (cedula, observacion):
            celda.alignment = Alignment(vertical="top", wrap_text=celda is observacion)
            celda.border = BORDE
            if fila.errores or fila.avisos:
                celda.fill = ROJO
                celda.font = Font(color="9C0006")
            elif fila.fallas:
                celda.fill = AMARILLO
                celda.font = Font(color="9C5700")
        if fila.avisos and not fila.errores:
            nota = "\n".join(a[:1].upper() + a[1:] for a in fila.avisos)
            observacion.comment = Comment(nota, NOMBRE, width=320, height=110)
    hoja.column_dimensions["A"].width = 15
    hoja.column_dimensions["B"].width = 120
    hoja.freeze_panes = "A2"
    if ruta is None:
        SALIDAS.mkdir(exist_ok=True)
        base = f"Observaciones {datetime.now():%Y-%m-%d %H.%M}"
        ruta, contador = SALIDAS / f"{base}.xlsx", 2
        while ruta.exists():
            ruta, contador = SALIDAS / f"{base} ({contador}).xlsx", contador + 1
    libro.save(ruta)
    return ruta
